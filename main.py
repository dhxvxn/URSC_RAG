"""Offline notebook-based document QA service.

The service is deliberately self-contained: SQLite stores the document registry,
notebook membership, chunks, and embeddings. The only network call made while
answering a question is to the Ollama container on the local Docker network.
"""

from __future__ import annotations

import base64
import hashlib
import io
import os
import re
import sqlite3
import threading
import time
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, List

import httpx
import numpy as np
import pandas as pd
import pymupdf as fitz
import uvicorn
from docx import Document as DocxDocument
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from sentence_transformers import SentenceTransformer


APP_TITLE = "Offline Notebook Document QA"
DATA_DIR = Path(os.getenv("DATA_DIR", "/data"))
DB_PATH = Path(os.getenv("DB_PATH", str(DATA_DIR / "knowledge.sqlite3")))
DOCUMENT_DIR = DATA_DIR / "documents"
IMAGE_DIR = DATA_DIR / "images"
MODEL_DIR = Path(os.getenv("EMBEDDING_MODEL_DIR", str(DATA_DIR / "models" / "embeddings")))
OCR_MODEL_DIR = Path(os.getenv("EASYOCR_MODULE_PATH", str(DATA_DIR / "models" / "easyocr")))
EMBEDDING_MODEL_NAME = os.getenv("EMBEDDING_MODEL", "all-MiniLM-L6-v2")
OLLAMA_BASE_URL = os.getenv("OLLAMA_BASE_URL", "http://ollama:11434").rstrip("/")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "llama3.2:3b")
OLLAMA_VISION_MODEL = os.getenv("OLLAMA_VISION_MODEL", "llava:7b")
OLLAMA_KEEP_ALIVE = os.getenv("OLLAMA_KEEP_ALIVE", "-1")
TOP_K_DEFAULT = int(os.getenv("TOP_K_DEFAULT", "6"))
MAX_CONTEXT_CHARS = int(os.getenv("MAX_CONTEXT_CHARS", "12000"))
MAX_VISION_IMAGES = int(os.getenv("MAX_VISION_IMAGES", "2"))
OLLAMA_NUM_CTX = int(os.getenv("OLLAMA_NUM_CTX", "4096"))
OLLAMA_NUM_PREDICT = int(os.getenv("OLLAMA_NUM_PREDICT", "512"))

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xls", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
DB_LOCK = threading.RLock()
_embedding_model: SentenceTransformer | None = None
_ocr_reader: Any = None
_ocr_lock = threading.Lock()
_retrieval_cache: dict[str, tuple[int, np.ndarray, list[dict[str, Any]]]] = {}
_retrieval_generation = 0
_retrieval_cache_lock = threading.RLock()


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def ensure_dirs() -> None:
    for directory in (DATA_DIR, DOCUMENT_DIR, IMAGE_DIR, DB_PATH.parent, MODEL_DIR, OCR_MODEL_DIR):
        directory.mkdir(parents=True, exist_ok=True)


def db() -> sqlite3.Connection:
    connection = sqlite3.connect(DB_PATH, timeout=60)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute("PRAGMA foreign_keys=ON")
    return connection


def init_db() -> None:
    ensure_dirs()
    with DB_LOCK, db() as connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS notebooks (
                id TEXT PRIMARY KEY,
                name TEXT NOT NULL UNIQUE,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS documents (
                id TEXT PRIMARY KEY,
                sha256 TEXT NOT NULL UNIQUE,
                filename TEXT NOT NULL,
                extension TEXT NOT NULL,
                stored_path TEXT NOT NULL,
                size_bytes INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                indexed_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS notebook_documents (
                notebook_id TEXT NOT NULL REFERENCES notebooks(id) ON DELETE CASCADE,
                document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                source_name TEXT NOT NULL,
                added_at TEXT NOT NULL,
                PRIMARY KEY (notebook_id, document_id)
            );
            CREATE TABLE IF NOT EXISTS chunks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                page INTEGER NOT NULL DEFAULT 0,
                kind TEXT NOT NULL DEFAULT 'text',
                text TEXT NOT NULL,
                embedding BLOB NOT NULL,
                embedding_dim INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS images (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                page INTEGER NOT NULL DEFAULT 0,
                stored_path TEXT NOT NULL,
                mime_type TEXT NOT NULL,
                ocr_text TEXT NOT NULL DEFAULT ''
            );
            CREATE INDEX IF NOT EXISTS idx_chunks_document ON chunks(document_id);
            CREATE INDEX IF NOT EXISTS idx_membership_notebook ON notebook_documents(notebook_id);
            CREATE INDEX IF NOT EXISTS idx_images_document ON images(document_id);
            """
        )
        # FTS5 provides a fast lexical signal for names, acronyms, identifiers,
        # and exact terminology that semantic embeddings can underweight.
        try:
            connection.execute(
                "CREATE VIRTUAL TABLE IF NOT EXISTS chunk_fts USING fts5(text, content='chunks', content_rowid='id')"
            )
            connection.executescript(
                """
                CREATE TRIGGER IF NOT EXISTS chunks_ai AFTER INSERT ON chunks BEGIN
                    INSERT INTO chunk_fts(rowid, text) VALUES (new.id, new.text);
                END;
                CREATE TRIGGER IF NOT EXISTS chunks_ad AFTER DELETE ON chunks BEGIN
                    INSERT INTO chunk_fts(chunk_fts, rowid, text) VALUES ('delete', old.id, old.text);
                END;
                CREATE TRIGGER IF NOT EXISTS chunks_au AFTER UPDATE ON chunks BEGIN
                    INSERT INTO chunk_fts(chunk_fts, rowid, text) VALUES ('delete', old.id, old.text);
                    INSERT INTO chunk_fts(rowid, text) VALUES (new.id, new.text);
                END;
                """
            )
            fts_count = connection.execute("SELECT COUNT(*) FROM chunk_fts").fetchone()[0]
            chunk_count = connection.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
            if fts_count != chunk_count:
                connection.execute("INSERT INTO chunk_fts(chunk_fts) VALUES ('rebuild')")
        except sqlite3.OperationalError:
            # Python builds without FTS5 still work with semantic retrieval.
            pass


def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def get_embedding_model() -> SentenceTransformer:
    global _embedding_model
    if _embedding_model is None:
        _embedding_model = SentenceTransformer(
            str(MODEL_DIR) if MODEL_DIR.exists() and any(MODEL_DIR.iterdir()) else EMBEDDING_MODEL_NAME,
            cache_folder=str(MODEL_DIR),
        )
    return _embedding_model


def get_ocr_reader() -> Any:
    global _ocr_reader
    if _ocr_reader is None:
        with _ocr_lock:
            if _ocr_reader is None:
                try:
                    import easyocr

                    _ocr_reader = easyocr.Reader(
                        ["en"],
                        model_storage_directory=str(OCR_MODEL_DIR),
                        download_enabled=False,
                        verbose=False,
                    )
                except Exception:
                    _ocr_reader = False
    return _ocr_reader


def ocr_array(array: np.ndarray) -> str:
    reader = get_ocr_reader()
    if not reader:
        return ""
    try:
        return " ".join(str(value) for value in reader.readtext(array, detail=0)).strip()
    except Exception:
        return ""


def image_array(image_bytes: bytes) -> np.ndarray:
    from PIL import Image

    with Image.open(io.BytesIO(image_bytes)) as image:
        return np.asarray(image.convert("RGB"))


def image_mime(extension: str) -> str:
    return {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".tif": "image/tiff",
        ".tiff": "image/tiff",
        ".bmp": "image/bmp",
        ".webp": "image/webp",
    }.get(extension, "image/png")


def split_text(text: str, max_chars: int = 1400, overlap: int = 180) -> list[str]:
    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()
    if not cleaned:
        return []
    if len(cleaned) <= max_chars:
        return [cleaned]
    pieces: list[str] = []
    start = 0
    while start < len(cleaned):
        end = min(len(cleaned), start + max_chars)
        if end < len(cleaned):
            boundary = cleaned.rfind(" ", start + max_chars // 2, end)
            if boundary > start:
                end = boundary
        pieces.append(cleaned[start:end].strip())
        if end >= len(cleaned):
            break
        start = max(0, end - overlap)
    return pieces


def extract_pdf(content: bytes) -> tuple[list[tuple[int, str]], list[tuple[int, bytes, str, str]]]:
    sections: list[tuple[int, str]] = []
    images: list[tuple[int, bytes, str, str]] = []
    document = fitz.open(stream=content, filetype="pdf")
    for page_index, page in enumerate(document, start=1):
        page_text = page.get_text("text").strip()
        if not page_text:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            array = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(pixmap.height, pixmap.width, pixmap.n)
            page_text = ocr_array(array)
        image_ocr: list[str] = []
        for image_index, image_info in enumerate(page.get_images(full=True), start=1):
            try:
                extracted = document.extract_image(image_info[0])
                image_bytes = extracted["image"]
                extension = "." + extracted.get("ext", "png")
                text = ocr_array(image_array(image_bytes))
                images.append((page_index, image_bytes, extension, text))
                if text:
                    image_ocr.append(f"[Image {image_index} OCR] {text}")
            except Exception:
                continue
        if image_ocr:
            page_text = (page_text + "\n" if page_text else "") + "\n".join(image_ocr)
        if page_text:
            sections.append((page_index, page_text))
    return sections, images


def extract_docx(content: bytes) -> tuple[list[tuple[int, str]], list[tuple[int, bytes, str, str]]]:
    document = DocxDocument(io.BytesIO(content))
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    images: list[tuple[int, bytes, str, str]] = []
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for name in archive.namelist():
            if name.startswith("word/media/"):
                extension = Path(name).suffix.lower() or ".png"
                image_bytes = archive.read(name)
                images.append((0, image_bytes, extension, ocr_array(image_array(image_bytes))))
    image_text = "\n".join(f"[Embedded image OCR] {item[3]}" for item in images if item[3])
    if image_text:
        lines.append(image_text)
    return ([(0, "\n".join(lines))] if lines else []), images


def extract_excel(content: bytes) -> tuple[list[tuple[int, str]], list[tuple[int, bytes, str, str]]]:
    sections: list[tuple[int, str]] = []
    sheets = pd.read_excel(io.BytesIO(content), sheet_name=None)
    for sheet_number, (sheet_name, frame) in enumerate(sheets.items(), start=1):
        sections.append((sheet_number, f"[Excel sheet: {sheet_name}]\n{frame.to_string(index=False)}"))
    images: list[tuple[int, bytes, str, str]] = []
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for name in archive.namelist():
            if name.startswith("xl/media/"):
                extension = Path(name).suffix.lower() or ".png"
                image_bytes = archive.read(name)
                images.append((0, image_bytes, extension, ocr_array(image_array(image_bytes))))
    if images:
        extra = "\n".join(f"[Embedded spreadsheet image OCR] {item[3]}" for item in images if item[3])
        if extra:
            sections.append((0, extra))
    return sections, images


def extract_document(content: bytes, filename: str) -> tuple[list[tuple[int, str]], list[tuple[int, bytes, str, str]]]:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        return extract_pdf(content)
    if extension == ".docx":
        return extract_docx(content)
    if extension in {".xlsx", ".xls"}:
        return extract_excel(content)
    if extension in IMAGE_EXTENSIONS:
        text = ocr_array(image_array(content))
        return ([(1, text)] if text else []), [(1, content, extension, text)]
    raise ValueError(f"Unsupported file type: {extension}")


def embed_texts(texts: list[str]) -> np.ndarray:
    if not texts:
        return np.empty((0, 0), dtype=np.float32)
    model = get_embedding_model()
    vectors = model.encode(texts, normalize_embeddings=True, convert_to_numpy=True, show_progress_bar=False)
    return np.asarray(vectors, dtype=np.float32)


def invalidate_retrieval_cache() -> None:
    global _retrieval_generation
    with _retrieval_cache_lock:
        _retrieval_generation += 1
        _retrieval_cache.clear()


def notebook_index(notebook_id: str) -> tuple[np.ndarray, list[dict[str, Any]]]:
    """Load one notebook's embeddings once and reuse them for later questions."""
    with _retrieval_cache_lock:
        cached = _retrieval_cache.get(notebook_id)
        if cached and cached[0] == _retrieval_generation:
            return cached[1], cached[2]
    with db() as connection:
        rows = connection.execute(
            """
            SELECT c.id, c.document_id, c.page, c.text, c.embedding, d.filename
            FROM chunks c
            JOIN notebook_documents nd ON nd.document_id = c.document_id
            JOIN documents d ON d.id = c.document_id
            WHERE nd.notebook_id = ? ORDER BY c.id
            """,
            (notebook_id,),
        ).fetchall()
    metadata = [
        {"id": row["id"], "document_id": row["document_id"], "filename": row["filename"], "page": row["page"], "text": row["text"]}
        for row in rows
    ]
    if rows:
        matrix = np.vstack([np.frombuffer(row["embedding"], dtype=np.float32) for row in rows]).astype(np.float32, copy=False)
    else:
        matrix = np.empty((0, 0), dtype=np.float32)
    with _retrieval_cache_lock:
        _retrieval_cache[notebook_id] = (_retrieval_generation, matrix, metadata)
    return matrix, metadata


def lexical_candidates(notebook_id: str, question: str, limit: int) -> list[int]:
    terms = re.findall(r"[A-Za-z0-9_]{2,}", question.lower())
    if not terms:
        return []
    match_expression = " OR ".join(f'"{term.replace(chr(34), "")}"' for term in terms[:24])
    try:
        with db() as connection:
            rows = connection.execute(
                """
                SELECT c.id
                FROM chunk_fts
                JOIN chunks c ON c.id = chunk_fts.rowid
                JOIN notebook_documents nd ON nd.document_id = c.document_id
                WHERE nd.notebook_id = ? AND chunk_fts MATCH ?
                ORDER BY bm25(chunk_fts) LIMIT ?
                """,
                (notebook_id, match_expression, limit),
            ).fetchall()
        return [int(row["id"]) for row in rows]
    except sqlite3.OperationalError:
        return []


def notebook_exists(notebook_id: str) -> bool:
    with db() as connection:
        return connection.execute("SELECT 1 FROM notebooks WHERE id = ?", (notebook_id,)).fetchone() is not None


def create_notebook_record(name: str) -> dict[str, str]:
    clean_name = name.strip()
    if not clean_name:
        raise ValueError("Notebook name cannot be empty")
    notebook_id = uuid.uuid4().hex
    try:
        with DB_LOCK, db() as connection:
            connection.execute("INSERT INTO notebooks(id, name, created_at) VALUES (?, ?, ?)", (notebook_id, clean_name, utc_now()))
    except sqlite3.IntegrityError as error:
        raise ValueError("A notebook with that name already exists") from error
    return {"id": notebook_id, "name": clean_name}


def ingest_document(notebook_id: str, filename: str, content: bytes) -> dict[str, Any]:
    if not notebook_exists(notebook_id):
        raise HTTPException(status_code=404, detail="Notebook not found")
    safe_filename = Path(filename.replace("\\", "/")).name or "document"
    extension = Path(safe_filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        return {"filename": safe_filename, "status": "skipped", "reason": "unsupported extension"}
    digest = sha256_bytes(content)
    with DB_LOCK, db() as connection:
        existing = connection.execute("SELECT id FROM documents WHERE sha256 = ?", (digest,)).fetchone()
        if existing:
            connection.execute("INSERT OR IGNORE INTO notebook_documents(notebook_id, document_id, source_name, added_at) VALUES (?, ?, ?, ?)", (notebook_id, existing["id"], safe_filename, utc_now()))
            invalidate_retrieval_cache()
            return {"filename": safe_filename, "status": "cached", "document_id": existing["id"]}
    try:
        sections, extracted_images = extract_document(content, safe_filename)
        flattened = [(page, piece) for page, text in sections for piece in split_text(text)]
        if not flattened and extracted_images:
            flattened = [(page, "[Document image; use the stored image source for visual details]") for page, *_ in extracted_images]
        vectors = embed_texts([piece for _, piece in flattened])
    except Exception as error:
        return {"filename": safe_filename, "status": "error", "reason": str(error)}
    if not flattened:
        return {"filename": safe_filename, "status": "error", "reason": "No readable text or OCR content found"}

    document_id = uuid.uuid4().hex
    stored_relative = Path("documents") / f"{digest}{extension}"
    stored_path = DATA_DIR / stored_relative
    stored_path.parent.mkdir(parents=True, exist_ok=True)
    stored_path.write_bytes(content)
    image_rows: list[tuple[int, str, str, str]] = []
    for image_index, (page, image_bytes, image_extension, ocr_text) in enumerate(extracted_images, start=1):
        relative_image = Path("images") / document_id / f"page-{page or 0}-{image_index}{image_extension}"
        absolute_image = DATA_DIR / relative_image
        absolute_image.parent.mkdir(parents=True, exist_ok=True)
        absolute_image.write_bytes(image_bytes)
        image_rows.append((page, relative_image.as_posix(), image_mime(image_extension), ocr_text))

    with DB_LOCK, db() as connection:
        connection.execute("INSERT INTO documents(id, sha256, filename, extension, stored_path, size_bytes, created_at, indexed_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", (document_id, digest, safe_filename, extension, stored_relative.as_posix(), len(content), utc_now(), utc_now()))
        for (page, text), vector in zip(flattened, vectors):
            connection.execute("INSERT INTO chunks(document_id, page, kind, text, embedding, embedding_dim) VALUES (?, ?, 'text', ?, ?, ?)", (document_id, page, text, vector.tobytes(), int(vector.shape[0])))
        for page, relative_path, mime_type, ocr_text in image_rows:
            connection.execute("INSERT INTO images(document_id, page, stored_path, mime_type, ocr_text) VALUES (?, ?, ?, ?, ?)", (document_id, page, relative_path, mime_type, ocr_text))
        connection.execute("INSERT INTO notebook_documents(notebook_id, document_id, source_name, added_at) VALUES (?, ?, ?, ?)", (notebook_id, document_id, safe_filename, utc_now()))
    invalidate_retrieval_cache()
    return {"filename": safe_filename, "status": "indexed", "document_id": document_id, "chunks": len(flattened), "images": len(image_rows)}


def retrieve(notebook_id: str, question: str, top_k: int) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not notebook_exists(notebook_id):
        raise HTTPException(status_code=404, detail="Notebook not found")
    query_vector = embed_texts([question])[0]
    matrix, metadata = notebook_index(notebook_id)
    if not metadata:
        return [], []
    candidate_limit = max(top_k * 8, 40)
    semantic_scores = matrix @ query_vector
    semantic_order = np.argsort(-semantic_scores)[: min(candidate_limit, len(metadata))]
    lexical_ids = lexical_candidates(notebook_id, question, candidate_limit)
    metadata_by_id = {item["id"]: item for item in metadata}
    fused_scores: dict[int, float] = {}
    semantic_values: dict[int, float] = {}
    # Reciprocal-rank fusion makes exact terms and semantic similarity
    # complementary without requiring a collection-specific score threshold.
    for rank, index in enumerate(semantic_order, start=1):
        chunk_id = metadata[index]["id"]
        fused_scores[chunk_id] = fused_scores.get(chunk_id, 0.0) + 1.0 / (60.0 + rank)
        semantic_values[chunk_id] = float(semantic_scores[index])
    for rank, chunk_id in enumerate(lexical_ids, start=1):
        fused_scores[chunk_id] = fused_scores.get(chunk_id, 0.0) + 1.0 / (60.0 + rank)
    selected_ids = sorted(fused_scores, key=lambda chunk_id: fused_scores[chunk_id], reverse=True)[: max(1, min(top_k, 20))]
    hits = [
        {**metadata_by_id[chunk_id], "score": round(semantic_values.get(chunk_id, 0.0), 4), "hybrid_score": round(fused_scores[chunk_id], 6)}
        for chunk_id in selected_ids
    ]
    with db() as connection:
        images: list[dict[str, Any]] = []
        for document_id in {item["document_id"] for item in hits}:
            image_rows = connection.execute("SELECT document_id, page, stored_path, mime_type, ocr_text FROM images WHERE document_id = ? LIMIT 6", (document_id,)).fetchall()
            images.extend(dict(row) for row in image_rows)
    return hits, images[:12]


def extractive_fallback(hits: list[dict[str, Any]]) -> str:
    if not hits:
        return "No indexed passages matched this notebook."
    lines = ["Ollama is unavailable; these are the most relevant indexed passages:"]
    for hit in hits[:5]:
        lines.append(f"\n[{hit['filename']}, page {hit['page'] or 'n/a'}]\n{hit['text']}")
    return "\n".join(lines)


async def ask_ollama(question: str, hits: list[dict[str, Any]], images: list[dict[str, Any]]) -> str | None:
    if not hits:
        return "No indexed passages matched this notebook."
    context = "\n\n".join(f"[{hit['filename']} page {hit['page'] or 'n/a'}]\n{hit['text']}" for hit in hits)
    prompt = ("Answer using only the supplied notebook sources. Cite the filename and page when available. "
              "If the sources do not contain the answer, say so. Describe relevant images or diagrams when they help.\n\n"
              f"Sources:\n{context[:MAX_CONTEXT_CHARS]}\n\nQuestion: {question}")
    image_payload: list[str] = []
    for image in images[:MAX_VISION_IMAGES]:
        try:
            image_payload.append(base64.b64encode((DATA_DIR / image["stored_path"]).read_bytes()).decode("ascii"))
        except OSError:
            continue
    model = OLLAMA_VISION_MODEL if image_payload else OLLAMA_MODEL
    message: dict[str, Any] = {"role": "user", "content": prompt}
    if image_payload:
        message["images"] = image_payload
    try:
        async with httpx.AsyncClient(timeout=180) as client:
            response = await client.post(
                f"{OLLAMA_BASE_URL}/api/chat",
                json={
                    "model": model,
                    "messages": [{"role": "system", "content": "You are an offline document research assistant."}, message],
                    "stream": False,
                    "keep_alive": OLLAMA_KEEP_ALIVE,
                    "options": {"num_ctx": OLLAMA_NUM_CTX, "num_predict": OLLAMA_NUM_PREDICT, "temperature": 0.1},
                },
            )
            response.raise_for_status()
            answer = response.json().get("message", {}).get("content", "").strip()
            return answer or None
    except Exception:
        return None


class NotebookCreate(BaseModel):
    name: str = Field(min_length=1, max_length=120)


class AskRequest(BaseModel):
    question: str = Field(min_length=1)
    top_k: int = Field(default=TOP_K_DEFAULT, ge=1, le=20)


app = FastAPI(title=APP_TITLE, version="2.0.0")
init_db()
app.mount("/media", StaticFiles(directory=str(DATA_DIR)), name="media")


@app.on_event("startup")
async def warm_up() -> None:
    # Loading the small embedding model once at startup removes the first-query
    # penalty. A missing model is reported on the first indexing request.
    if os.getenv("PRELOAD_EMBEDDING_MODEL", "1").lower() not in {"0", "false", "no"}:
        try:
            get_embedding_model()
        except Exception:
            pass


@app.get("/health")
def health() -> dict[str, Any]:
    with db() as connection:
        notebook_count = connection.execute("SELECT COUNT(*) FROM notebooks").fetchone()[0]
        document_count = connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
        chunk_count = connection.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
    return {"status": "ok", "notebooks": notebook_count, "documents": document_count, "chunks": chunk_count, "ollama": OLLAMA_BASE_URL, "retrieval": "hybrid in-memory embeddings + SQLite FTS5"}


@app.get("/api/notebooks")
def list_notebooks() -> list[dict[str, Any]]:
    with db() as connection:
        rows = connection.execute("SELECT n.id, n.name, n.created_at, COUNT(nd.document_id) AS documents FROM notebooks n LEFT JOIN notebook_documents nd ON nd.notebook_id = n.id GROUP BY n.id ORDER BY n.created_at").fetchall()
    return [dict(row) for row in rows]


@app.post("/api/notebooks")
def create_notebook(payload: NotebookCreate) -> dict[str, str]:
    try:
        return create_notebook_record(payload.name)
    except ValueError as error:
        raise HTTPException(status_code=409, detail=str(error)) from error


@app.get("/api/notebooks/{notebook_id}/documents")
def list_documents(notebook_id: str) -> list[dict[str, Any]]:
    if not notebook_exists(notebook_id):
        raise HTTPException(status_code=404, detail="Notebook not found")
    with db() as connection:
        rows = connection.execute("SELECT d.id, nd.source_name, d.filename, d.size_bytes, d.indexed_at, COUNT(c.id) AS chunks, COUNT(DISTINCT i.id) AS images FROM notebook_documents nd JOIN documents d ON d.id = nd.document_id LEFT JOIN chunks c ON c.document_id = d.id LEFT JOIN images i ON i.document_id = d.id WHERE nd.notebook_id = ? GROUP BY d.id ORDER BY nd.added_at DESC", (notebook_id,)).fetchall()
    return [dict(row) for row in rows]


@app.post("/api/notebooks/{notebook_id}/documents")
async def upload_documents(notebook_id: str, files: List[UploadFile] = File(...)) -> dict[str, Any]:
    if not notebook_exists(notebook_id):
        raise HTTPException(status_code=404, detail="Notebook not found")
    results = []
    for upload in files:
        results.append(ingest_document(notebook_id, upload.filename or "document", await upload.read()))
    return {"notebook_id": notebook_id, "results": results}


@app.post("/api/notebooks/{notebook_id}/ask")
async def ask_notebook(notebook_id: str, request: AskRequest) -> dict[str, Any]:
    started = time.perf_counter()
    hits, images = retrieve(notebook_id, request.question, request.top_k)
    retrieved_ms = round((time.perf_counter() - started) * 1000, 1)
    answer_started = time.perf_counter()
    answer = await ask_ollama(request.question, hits, images)
    return {"notebook_id": notebook_id, "answer": answer or extractive_fallback(hits), "sources": [{**hit, "image_urls": [f"/media/{image['stored_path']}" for image in images if image["document_id"] == hit["document_id"] and (not hit["page"] or image["page"] == hit["page"])]} for hit in hits], "images": [{**image, "url": f"/media/{image['stored_path']}"} for image in images], "retrieval": "hybrid in-memory embeddings + SQLite FTS5", "timings_ms": {"retrieval": retrieved_ms, "answer": round((time.perf_counter() - answer_started) * 1000, 1)}}


@app.get("/", response_class=HTMLResponse)
def web_interface() -> str:
    return """<!doctype html>
<html><head><meta charset='utf-8'><title>Offline Notebook QA</title>
<style>body{font-family:system-ui;max-width:1000px;margin:2rem auto;padding:0 1rem;background:#f4f6f8}main{background:white;padding:1.5rem;border-radius:12px;box-shadow:0 2px 12px #0001}label{display:block;margin-top:1rem;font-weight:600}input,select,textarea,button{width:100%;box-sizing:border-box;padding:.7rem;margin-top:.35rem}button{cursor:pointer;background:#165dff;color:white;border:0;border-radius:6px}.row{display:grid;grid-template-columns:1fr 1fr;gap:1rem}#answer{white-space:pre-wrap;background:#f0f7ff;padding:1rem;margin-top:1rem;border-radius:8px}.thumb{max-width:300px;max-height:200px;margin:.5rem;border-radius:5px}</style></head>
<body><main><h1>Offline Notebook QA</h1><p>Documents are embedded once and retained in the local data volume.</p>
<div class='row'><div><label>Notebook</label><select id='notebook'></select></div><div><label>New notebook</label><input id='newName' placeholder='e.g. Propulsion'></div></div>
<button onclick='createNotebook()'>Create notebook</button><label>Add files or a folder</label><input id='files' type='file' multiple webkitdirectory directory><button onclick='upload()'>Index selected files</button>
<label>Question</label><textarea id='question' rows='4' placeholder='Ask about this notebook'></textarea><button onclick='ask()'>Ask</button><div id='status'></div><div id='answer'></div></main>
<script>
const $=id=>document.getElementById(id); let notebooks=[];
async function refresh(){notebooks=await (await fetch('/api/notebooks')).json();$('notebook').innerHTML=notebooks.map(n=>`<option value="${n.id}">${n.name} (${n.documents} files)</option>`).join('');}
async function createNotebook(){let name=$('newName').value.trim();if(!name)return;let r=await fetch('/api/notebooks',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({name})});if(!r.ok){$('status').textContent=await r.text();return;}$('newName').value='';await refresh();}
async function upload(){let files=[...$('files').files];let id=$('notebook').value;if(!id||!files.length){$('status').textContent='Create/select a notebook and choose files.';return;}let form=new FormData();files.forEach(f=>form.append('files',f,f.webkitRelativePath||f.name));$('status').textContent='Indexing; repeated files are reused from the cache...';let r=await fetch(`/api/notebooks/${id}/documents`,{method:'POST',body:form});$('status').textContent=await r.text();await refresh();}
async function ask(){let id=$('notebook').value,q=$('question').value.trim();if(!id||!q)return;$('answer').textContent='Retrieving...';let r=await fetch(`/api/notebooks/${id}/ask`,{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({question:q})});let d=await r.json();$('answer').textContent=d.answer+'\n\nSources:\n'+d.sources.map(s=>`${s.filename} (page ${s.page||'n/a'})\n${s.text}`).join('\n\n');if(d.images&&d.images.length){let heading=document.createElement('div');heading.textContent='\n\nRetrieved images:';$('answer').appendChild(heading);d.images.forEach(image=>{let img=document.createElement('img');img.className='thumb';img.alt=image.ocr_text||'Retrieved document image';img.src=image.url;$('answer').appendChild(img);});}}
refresh();
</script></body></html>"""


if __name__ == "__main__":
    uvicorn.run("main:app", host="0.0.0.0", port=8000)
